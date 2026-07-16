# -*- coding: utf-8 -*-
"""
สคริปต์นี้ทำงานสมบูรณ์แล้ว — ห้ามเขียนโค้ดส่วน RENDERING ใหม่เด็ดขาด
ถ้าต้องการสร้างเอกสารชุดใหม่ ให้แก้ไขเฉพาะใน DATA = {...} ด้านล่างนี้เท่านั้น
แล้วรัน:  python generate_boq.py
ผลลัพธ์: ไฟล์ .xlsx พร้อมพิมพ์ A4 (เปิดใน Excel แล้วกด Print ได้เลย ไม่ต้องแปลง PDF)
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

# #########################################################################
# ██  DATA — แก้ตรงนี้ที่เดียวเท่านั้น เพื่อสร้างเอกสารชุดใหม่  ██
# #########################################################################
DATA = {
    "output_filename": "output.xlsx",

    "title_line1": "ใบถอดต้นทุน EV CHARGING STATION",
    "title_line2": "MPJ LOGISTIC  |  สถานที่: ศรีราชา ชลบุรี",

    "doc_no": "PP69001",
    "salesperson": "Sittipong",
    "date_text": "15 ก.ค. 2569",

    "charger_name": "Group Charger: 720 kW",
    "charger_qty": "1 เครื่อง",

    # แต่ละ section: title = หัวข้อแถบสีน้ำเงิน, is_extra = True สำหรับ section
    # แบบ "อุปกรณ์และเงื่อนไขเพิ่มเติม" (ใช้ฟอนต์เล็กกว่า header 5pt/body 6pt),
    # False สำหรับ section ทั่วไป (header 7pt/body 6pt)
    # items แต่ละแถว = (รหัส, ประเภท, รายการสินค้า, จำนวน, ระยะ(m))
    # ถ้ารหัส == "-" แถวนั้นจะเรนเดอร์เป็นตัวเอียงสีเทาอัตโนมัติ (รายการย่อย/หมายเหตุ)
    "sections": [
        {
            "title": "1. ระบบแรงสูง (High Voltage System)",
            "is_extra": False,
            "items": [
                ("7.3.11", "หม้อแปลง 22 (24) kV / 416 V", "1000 kVA", "1", ""),
                ("4.1", "ระบบแรงสูง", "ชุดรับไฟแรงสูง 22 ( 24 kV. )", "1", ""),
                ("-", "-", "- ติดตั้งนั่งร้านหม้อแปลง 2 ต้น", "1", ""),
                ("-", "-", "- ต้นรับไฟจากการไฟฟ้า 1 ต้น", "1", ""),
                ("6", "ระบบแรงสูง", "ชุดสายไฟแรงสูง 22 ( 24 kV.) 3P", "1", "20"),
            ],
        },
        {
            "title": "2. ระบบแรงต่ำ (Low Voltage System)",
            "is_extra": False,
            "items": [
                ("9.12.34", "(TR to MDB) TR to Land - สาย 3P 4W ร้อยท่อ กลุ่ม 5 ฝังใต้ดิน", "CV 5 ( 4 x 150 ) / 5 x 95 THWG, 5 x 90 มม.", "1", "9 m.(9 m.)"),
                ("9.12.34", "(TR to MDB) Land to MDB - สาย 3P 4W ร้อยท่อ กลุ่ม 5 ฝังใต้ดิน", "CV 5 ( 4 x 150 ) / 5 x 95 THWG, 5 x 90 มม.", "1", "13 m.(10 m.)"),
                ("-", "Main MCCB", "ABB (1200 A (AT)/1200 A (AF))", "1", ""),
                ("4.70.10", "Sub MCCB", "ABB MCCB SUB C1 450 A (3 ชุด)", "3", ""),
                ("ตู้MDB", "ตู้ MDB", "120 x 210 x 80 (MDB + busbar)", "1", ""),
                ("ตู้MDB", "ตู้ MDB", "ระบบ Ground MDB", "1", ""),
                ("9.12.19", "Charger 1: สาย 3P 4W ร้อยท่อ กลุ่ม 5 ฝังใต้ดิน", "CV 3 SET OF 2 ( 4 x 120 ) / 2 x 70 THWG, HDPE 3 ( 2 x 90 ) มม.", "1", "8 m.(5 m.)"),
                ("-", "Terminal 1 Cable config (สาย 3P 4W ร้อยท่อ กลุ่ม 5 ฝังใต้ดิน)", "(350A) - CV 2 ( 2 x 185 ) / 2 x 95 THW, HDPE 2 x 75 มม.", "1", "10.5 m.(8 m.)"),
                ("-", "Terminal 2 Cable config (สาย 3P 4W ร้อยท่อ กลุ่ม 5 ฝังใต้ดิน)", "(350A) - CV 2 ( 2 x 185 ) / 2 x 95 THW, HDPE 2 x 75 มม.", "1", "18.5 m.(16 m.)"),
                ("-", "Terminal 3 Cable config (สาย 3P 4W ร้อยท่อ กลุ่ม 5 ฝังใต้ดิน)", "(350A) - CV 2 ( 2 x 185 ) / 2 x 95 THW, HDPE 2 x 75 มม.", "1", "24.5 m.(22 m.)"),
                ("-", "Terminal 4 Cable config (สาย 3P 4W ร้อยท่อ กลุ่ม 5 ฝังใต้ดิน)", "(350A) - CV 2 ( 2 x 185 ) / 2 x 95 THW, HDPE 2 x 75 มม.", "1", "34.5 m.(32 m.)"),
            ],
        },
        {
            "title": "3. อุปกรณ์และเงื่อนไขเพิ่มเติม (Additional Equipment)",
            "is_extra": True,
            "items": [
                ("5", "อุปกรณ์ประกอบสถานี", "เสากันชน (เหล็ก)", "4", ""),
                ("3", "อุปกรณ์ประกอบสถานี", "ขอบกั้นล้อ (ยาง)", "4", ""),
                ("2", "อุปกรณ์ประกอบสถานี", "ถังดับเพลิง+ตู้(CO2)", "1", ""),
                ("7", "อุปกรณ์ประกอบสถานี", "ป้ายสูง + วิธีใช้งาน", "1", ""),
                ("2", "ระบบสื่อสาร", "ROUTER 4G ใส่SIM + HUB 6 PORT + ตู้ติดตั้ง MDB", "1", ""),
                ("3", "ระบบสื่อสาร", "สาย LAN CAT6 + ท่อ PVC 16 มม.", "1", "20"),
                ("4", "ระบบสื่อสาร", "กล้องวงจรปิด + SIM", "4", ""),
                ("5", "ระบบสื่อสาร", "สาย AC กล้อง ( VAF 1.5 x 2 ) + ท่อ PVC 16 มม.", "1", "48"),
                ("6", "ระบบสื่อสาร", "หลอด LED 18 ม. PHILIPS T8", "24", ""),
                ("7", "ระบบสื่อสาร", "สาย AC หลอด ( 1.5 x 2 ) ท่อ PVC 16 มม.", "1", "31"),
                ("9.21.1", "งานปูน", "ฐานปูน MDB 1.5 ม. x 1.2 ม. x 0.2 ม.", "1", ""),
                ("9.21.2", "งานปูน", "ฐานปูน เครื่องชาร์จ 1.2 ม. x 0.8 ม. x 0.2 ม.", "4", ""),
                ("9.21.1", "งานปูน", "ฐานปูน Power Cabinet", "1", ""),
                ("9.24.8", "งานป้าย", "- ป้ายใช้งาน", "4", ""),
                ("9.24.9", "งานป้าย", "- ป้าย QR CODE 2 ชิ้น", "4", ""),
                ("-", "หลังคา", "หลังคา MDB", "1", ""),
                ("-", "หลังคา", "หลังคา CHARGER (ปกติ)", "4", ""),
            ],
        },
    ],

    "note": "หมายเหตุ: จำนวนและระยะทางเป็นค่าประมาณการเบื้องต้น อาจมีการเปลี่ยนแปลงตามหน้างานจริง",
}
# #########################################################################
# ██  จบส่วน DATA — ห้ามแก้โค้ดด้านล่างนี้  ██
# #########################################################################


# =========================================================================
# FONT — auto-detect a real installed Thai font FILE on this machine.
# We need the actual font file (not just a name string) to measure glyph
# widths for row-height calculation. Add your own path here if none match
# (Windows: right-click the font in C:\Windows\Fonts > Properties).
# =========================================================================
FONT_FILE_CANDIDATES = [
    r"C:\Windows\Fonts\THSarabunNew.ttf",
    r"C:\Windows\Fonts\THSarabun.ttf",
    r"C:\Windows\Fonts\tahoma.ttf",
    r"C:\Windows\Fonts\Tahoma.ttf",
    "/System/Library/Fonts/Supplemental/Tahoma.ttf",       # macOS
    "/Library/Fonts/THSarabunNew.ttf",                      # macOS (manual install)
    "/usr/share/fonts/truetype/thai/THSarabunNew.ttf",      # Linux
    "/usr/share/fonts/opentype/tlwg/Loma.otf",              # Linux (TLWG package)
]

def _find_font_file():
    for p in FONT_FILE_CANDIDATES:
        if os.path.exists(p):
            return p
    return None

FONT_FILE = _find_font_file()

if FONT_FILE and "THSarabun" in FONT_FILE:
    FONT_NAME = "TH Sarabun New"
elif FONT_FILE and "ahoma" in FONT_FILE:
    FONT_NAME = "Tahoma"
elif FONT_FILE and "Loma" in FONT_FILE:
    FONT_NAME = "Loma"
else:
    FONT_NAME = "TH Sarabun New"   # best-effort default even if unmeasurable here

print(f"[font] file={FONT_FILE!r} name={FONT_NAME!r}"
      + ("" if FONT_FILE else "  (WARNING: no font file found — row-height falls back to a conservative estimate. "
                               "Add your font's path to FONT_FILE_CANDIDATES for best accuracy.)"))

# ---- Colors ----
C_DARK, C_PRIMARY, C_PRIMARY_L = "1F3864", "2E5B9A", "DCE6F2"
C_ACCENT, C_ROW_ALT, C_BORDER = "1FA37A", "F4F7FB", "B9C6D8"
C_WHITE, C_TEXT, C_MUTED = "FFFFFF", "1A1A1A", "5B6B7F"
C_CHARGER_FILL = "E9F7F1"

thin = Side(style="thin", color=C_BORDER)
box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "รายการสินค้า"
ws.sheet_view.showGridLines = False

widths = {"A": 9, "B": 27, "C": 52, "D": 8, "E": 13}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
LAST_COL, LAST_COL_IDX = "E", 5

# ---- Font sizes (fixed spec — do not change) ----
SZ_HEADER, SZ_SECTION = 11, 9
SZ_TBL_HEAD, SZ_BODY = 7, 6
SZ_CHARGER = 8
SZ_EXTRA_HEAD, SZ_EXTRA_BODY = 5, 6
SZ_PAGE_NO = 7

# =========================================================================
# Text-wrap estimator — measures real glyph widths so row heights are
# always tall enough. DO NOT replace with a fixed/guessed row height.
# =========================================================================
try:
    from PIL import ImageFont
    _PIL_OK = FONT_FILE is not None
except ImportError:
    _PIL_OK = False

MDW, CELL_PAD_PT, SAFETY = 7, 6, 0.85
_font_cache = {}

def _get_font(size_pt):
    key = int(round(size_pt))
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(FONT_FILE, key)
    return _font_cache[key]

def _col_avail_width_pt(width_units, font_size, indent=1):
    px = int(width_units * MDW) + 5
    pt = px * 0.75
    indent_pt = indent * font_size * 0.6
    return max((pt - indent_pt - CELL_PAD_PT) * SAFETY, 10)

def est_lines(text, col_width_units, font_size=10, indent_chars=1):
    if not text:
        return 1
    if _PIL_OK:
        font = _get_font(font_size)
        max_w = _col_avail_width_pt(col_width_units, font_size, indent_chars)
        words, lines, cur = str(text).split(" "), 1, ""
        for w in words:
            trial = (cur + " " + w).strip()
            if font.getlength(trial) <= max_w or not cur:
                cur = trial
            else:
                lines += 1
                cur = w
        return lines
    else:
        chars_per_line = max(int(col_width_units * 1.0 * (10.0 / font_size)) - 2, 6)
        words, lines, cur = str(text).split(" "), 1, 0
        for w in words:
            wl = len(w) + 1
            if cur + wl > chars_per_line:
                lines += 1
                cur = wl
            else:
                cur += wl
        return lines

def row_height_for(cells, font_size=10, min_lines=1):
    line_pt, pad_pt = font_size * 1.8, font_size * 1.5
    max_lines = max([est_lines(t, widths[c], font_size) for t, c in cells] + [min_lines])
    return max_lines * line_pt + pad_pt

def set_row_height(r, h):
    ws.row_dimensions[r].height = h

def merge_and_write(cell_range, value, font=None, fill=None, align=None):
    ws.merge_cells(cell_range)
    c = ws[cell_range.split(":")[0]]
    c.value = value
    if font: c.font = font
    if fill: c.fill = fill
    if align: c.alignment = align
    return c

# =========================================================================
# RENDERING
# =========================================================================
row = 1
set_row_height(row, 8); row += 1

title_fill = PatternFill("solid", fgColor=C_DARK)
merge_and_write(f"A{row}:E{row}", DATA["title_line1"],
                 font=Font(name=FONT_NAME, size=SZ_HEADER, bold=True, color=C_WHITE),
                 fill=title_fill, align=Alignment(horizontal="left", vertical="center", indent=1))
set_row_height(row, 20); row += 1

merge_and_write(f"A{row}:E{row}", DATA["title_line2"],
                 font=Font(name=FONT_NAME, size=SZ_HEADER, bold=True, color=C_WHITE),
                 fill=title_fill, align=Alignment(horizontal="left", vertical="center", indent=1))
set_row_height(row, 18); row += 1

merge_and_write(f"A{row}:E{row}", "", fill=title_fill)
set_row_height(row, 5); row += 1
row += 1  # spacer

info_font = Font(name=FONT_NAME, size=SZ_HEADER, bold=True, color=C_TEXT)
info_fill = PatternFill("solid", fgColor=C_ROW_ALT)
info_row = row
set_row_height(info_row, 18)
merge_and_write(f"A{info_row}:B{info_row}", f"เอกสาร: {DATA['doc_no']}",
                 font=info_font, align=Alignment(horizontal="left", indent=1))
merge_and_write(f"C{info_row}:C{info_row}", f"พนักงานขาย: {DATA['salesperson']}",
                 font=info_font, align=Alignment(horizontal="left", indent=1))
merge_and_write(f"D{info_row}:E{info_row}", f"วันที่: {DATA['date_text']}",
                 font=info_font, align=Alignment(horizontal="right", indent=1))
for cc in range(1, LAST_COL_IDX + 1):
    ws.cell(row=info_row, column=cc).fill = info_fill
row += 1
set_row_height(row, 6); row += 1

HEADER_BLOCK_LAST_ROW = row - 1

def section_header(r, text, size=SZ_SECTION):
    fill = PatternFill("solid", fgColor=C_PRIMARY)
    merge_and_write(f"A{r}:E{r}", f"  {text}",
                     font=Font(name=FONT_NAME, size=size, bold=True, color=C_WHITE),
                     fill=fill, align=Alignment(horizontal="left", vertical="center", indent=1))
    set_row_height(r, size * 2.4 + 6)

def table_header(r, cols=("รหัส", "ประเภท", "รายการสินค้า", "จำนวน", "ระยะ (m)"), size=SZ_TBL_HEAD):
    fill = PatternFill("solid", fgColor=C_PRIMARY_L)
    f = Font(name=FONT_NAME, size=size, bold=True, color=C_DARK)
    for i, val in enumerate(cols):
        c = ws.cell(row=r, column=i + 1, value=val)
        c.font, c.fill, c.border = f, fill, box_border
        c.alignment = Alignment(horizontal="center" if i in (0, 3, 4) else "left",
                                 vertical="center", indent=0 if i in (0, 3, 4) else 1)
    set_row_height(r, size * 2.5 + 6)

def data_row(r, values, shaded=False, dash_row=False, size=SZ_BODY):
    fill = PatternFill("solid", fgColor=C_ROW_ALT) if shaded else PatternFill("solid", fgColor=C_WHITE)
    f = Font(name=FONT_NAME, size=size, color=C_MUTED if dash_row else C_TEXT, italic=dash_row)
    for i, val in enumerate(values):
        c = ws.cell(row=r, column=i + 1, value=val)
        c.fill, c.border, c.font = fill, box_border, f
        c.alignment = (Alignment(horizontal="center", vertical="center", wrap_text=True) if i in (0, 3, 4)
                        else Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    h = row_height_for([(values[1] if len(values) > 1 else "", "B"),
                         (values[2] if len(values) > 2 else "", "C")], font_size=size)
    set_row_height(r, h)

def spacer(r, h=8):
    set_row_height(r, h)

# ---- Charger callout ----
section_header(row, "เครื่องชาร์จ (Charger)"); row += 1
table_header(row, cols=("", "", "รายการสินค้า", "จำนวน", ""), size=SZ_CHARGER); row += 1
charger_fill = PatternFill("solid", fgColor=C_CHARGER_FILL)
c = ws.cell(row=row, column=3, value=DATA["charger_name"])
c.font, c.fill, c.border = Font(name=FONT_NAME, size=SZ_CHARGER, bold=True, color=C_DARK), charger_fill, box_border
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
d = ws.cell(row=row, column=4, value=DATA["charger_qty"])
d.font, d.fill, d.border = Font(name=FONT_NAME, size=SZ_CHARGER, bold=True, color=C_ACCENT), charger_fill, box_border
d.alignment = Alignment(horizontal="center", vertical="center")
for cc in (1, 2, 5):
    ws.cell(row=row, column=cc).fill = charger_fill
    ws.cell(row=row, column=cc).border = box_border
set_row_height(row, SZ_CHARGER * 2.8 + 5)
row += 1
spacer(row); row += 1

# ---- Sections (generic loop — driven entirely by DATA) ----
for section in DATA["sections"]:
    head_size = SZ_EXTRA_HEAD if section["is_extra"] else SZ_TBL_HEAD
    body_size = SZ_EXTRA_BODY if section["is_extra"] else SZ_BODY

    section_header(row, section["title"]); row += 1
    table_header(row, size=head_size); row += 1

    for i, item in enumerate(section["items"]):
        dash = item[0] == "-"
        data_row(row, item, shaded=(i % 2 == 1), dash_row=dash, size=body_size)
        row += 1

    spacer(row); row += 1

# ---- Footer note ----
note_fill = PatternFill("solid", fgColor=C_ROW_ALT)
merge_and_write(f"A{row}:E{row}", DATA["note"],
                 font=Font(name=FONT_NAME, size=7, italic=True, color=C_MUTED),
                 fill=note_fill, align=Alignment(horizontal="left", vertical="center", indent=1))
set_row_height(row, 16)
row += 1

LAST_ROW = row - 1

# =========================================================================
# PAGE SETUP — A4, print-ready straight from Excel (File > Print)
# =========================================================================
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.25)
ws.print_area = f"A1:{LAST_COL}{LAST_ROW}"
ws.print_options.horizontalCentered = True
ws.oddFooter.center.text = "หน้า &P จาก &N"
ws.oddFooter.center.size = SZ_PAGE_NO
ws.print_title_rows = f"1:{HEADER_BLOCK_LAST_ROW}"   # repeat header block on every printed page

wb.save(DATA["output_filename"])
print(f"[done] saved {DATA['output_filename']}  (last_row={LAST_ROW})")
print("[tip] If a section heading falls at the bottom of a page, open in Excel and use")
print("      Page Layout > Breaks > Insert Page Break before that section row.")
